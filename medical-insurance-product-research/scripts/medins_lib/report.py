"""生成分析报告（Markdown 必出；装有 python-docx 时同时输出 Word）。

体例：标题 → 目录 → 一、样本与主要结论 → 二、承保范围（核保端 / 理赔端）→ 三、保障形态与版本演进 →
四、创新点识别 → 五、趋势判断与风险提示 → 附录（证据清单、资料清单）。表题“表N　…”置于表上方，表注“注：…”置于表下方，
Word 版为三线表、正文首行缩进、页脚页码。

事实部分全部由标准化记录自动生成并附证据编号；研究者判断写在 <workspace>/analysis.json，
其中用 [[pid:字段]] 引用证据（报告里自动换成 [E#]），例如“免赔额由1万提高到1.5万[[p1a2b3c4d:deductible]]”。
没有 analysis.json 时，创新点等章节以“候选（待研究者判断）”形式输出，不会替研究者下结论。
"""
from __future__ import annotations

import re
from pathlib import Path

from .citations import build_index, fmt_refs, refs_for
from .common import MedinsError, read_json, today, truncate, write_text
from .doctypes import DOC_TYPES
from .fetch import completeness, load_manifest
from .records import field_map, load_record
from .workspace import Workspace

TOKEN = re.compile(r"\[\[([A-Za-z0-9_\-]+):([a-z0-9_]+)\]\]")
CAPTION = re.compile(r"^(附?表[A-Z]?\d+)　")
CN_NUM = "一二三四五六七八九十"

UNDERWRITING_KEYS = ["body_type", "access_channel", "hd_count", "hd_disease_scope", "hd_history_items", "hd_exemptions",
                     "hd_disclosure_effect", "nodules_three_highs", "entry_age", "underwriting_tools"]
CLAIM_KEYS = ["preexisting_definition", "preexisting_general", "preexisting_serious", "preexisting_special", "first_year_rules",
              "waiting_period"]
FORM_KEYS = ["term", "guaranteed_renewal", "renewal_rules", "deductible", "coinsurance", "no_social_insurance", "hospital_scope",
             "total_limit", "premium_50m", "premium_basis"]
BENEFIT_KEYS = ["general_medical", "critical_medical", "special_outpatient", "rehab", "outpatient", "external_drugs",
                "advanced_therapy", "other_benefits", "health_services"]


def _v(rec: dict, key: str) -> str:
    f = rec["fields"].get(key, {})
    st = f.get("status")
    if st in (None, "未填"):
        return "—"
    if st in ("待核", "不适用"):
        return f"〔{st}〕{f.get('note', '')}"
    tag = {"推定": "〔推定〕", "待复核": "〔待复核〕", "规则抽取": "〔待确认〕"}.get(st, "")
    return f"{f.get('value', '')}{tag}"


def _cell(rec, key, idx, pid, n=160):
    return f"{truncate(_v(rec, key), n)}{fmt_refs(refs_for(idx, pid, key))}"


def _md_table(headers: list[str], rows: list[list[str]]) -> str:
    def esc(x):
        return str(x if x is not None else "").replace("|", "／").replace("\n", "<br>")
    out = ["| " + " | ".join(esc(h) for h in headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def resolve_tokens(text: str, idx: dict, problems: list[str]) -> str:
    def rep(m):
        pid, field = m.group(1), m.group(2)
        ids = refs_for(idx, pid, field)
        if not ids:
            problems.append(f"analysis.json 引用无效或无证据：[[{pid}:{field}]]")
            return "[证据缺失]"
        return fmt_refs(ids)
    return TOKEN.sub(rep, text or "")


def check_claims(text: str, where: str, problems: list[str]) -> None:
    """含数字的判断句须带证据引用，否则提示。"""
    for sent in re.split(r"(?<=[。；！？\n])", text or ""):
        if re.search(r"\d", sent) and not TOKEN.search(sent) and "[E" not in sent and len(sent.strip()) > 8:
            problems.append(f"{where}：含数字的句子未附证据引用：「{truncate(sent.strip(), 50)}」")


def numbered(i: int, text: str) -> str:
    """“1．**首句。**其余”——首句（不超过 40 字）加粗，与研究报告的分点写法一致。"""
    text = (text or "").strip()
    m = re.match(r"^(.{4,40}?[。；：])(.+)$", text, re.S)
    return f"{i}．**{m.group(1)}**{m.group(2)}" if m else f"{i}．{text}"


def build_report(ws: Workspace, pids: list[str] | None = None, out_dir: Path | None = None, title: str = "",
                 docx: bool = True) -> dict:
    pids = [ws.get(p)["pid"] for p in (pids or list(ws.products()))]
    if not pids:
        raise MedinsError("工作区没有产品")
    prods = [ws.get(p) for p in pids]
    recs = {p["pid"]: load_record(ws, p["pid"]) for p in prods}
    fm = field_map()
    items, idx = build_index(ws, pids)
    cmp = read_json(ws.outputs_dir / "对比结果.json", {}) or {}
    inn = read_json(ws.outputs_dir / "创新点候选.json", {}) or {}
    val = read_json(ws.outputs_dir / "校验报告.json", {}) or {}
    analysis = read_json(ws.analysis_path, {}) or {}
    if not isinstance(analysis, dict):
        analysis = {}
    title = title or analysis.get("title") or (ws.load() or {}).get("title") or "医疗险产品研究"
    problems: list[str] = []
    secs = analysis.get("sections") or {}
    overall = analysis.get("overall") or {}
    counter = {"t": 0, "a": 0}

    def caption(text: str, appendix: bool = False) -> str:
        if appendix:
            counter["a"] += 1
            return f"附表{counter['a']}　{text}"
        counter["t"] += 1
        return f"表{counter['t']}　{text}"

    def say(key: str, where: str | None = None) -> list[str]:
        t = secs.get(key, "") if where is None else key
        if not (t or "").strip():
            return []
        check_claims(t, where or f"sections.{key}", problems)
        return [resolve_tokens(x.strip(), idx, problems) for x in t.split("\n") if x.strip()] + [""]

    def pname(p: dict) -> str:
        return f"{p['company']}·{p['product_name']}"

    def cmp_table(keys: list[str], n: int) -> str:
        if len(prods) <= 4:  # 维度为行、产品为列
            rows = [[fm[k]["label"]] + [_cell(recs[p["pid"]], k, idx, p["pid"], n) for p in prods] for k in keys]
            return _md_table(["维度"] + [pname(p) for p in prods], rows)
        rows = [[pname(p)] + [_cell(recs[p["pid"]], k, idx, p["pid"], max(40, n // 2)) for k in keys] for p in prods]
        return _md_table(["产品"] + [fm[k]["label"] for k in keys], rows)

    L: list[str] = [f"# {title}", "",
                    f"> 生成日期：{today()}　样本：{len(prods)} 款　事实性取值均取自已归档原件，[E#] 为证据编号（见附表1）；"
                    "〔推定〕为基于原文的归纳，〔待核〕为尚未取得资料。", ""]

    # 一、样本与主要结论
    L += ["## 一、样本与主要结论", "", "### （一）样本产品与资料来源", ""]
    L += say("scope")
    rows = []
    for i, p in enumerate(prods, 1):
        m = load_manifest(ws, p["pid"])
        comp = completeness(m)
        ok = sum(1 for v in comp.values() if v == "已归档" or "非缺文件" in v)
        lv = sorted({d.get("evidence_level", "?") for d in m["documents"]})
        r = recs[p["pid"]]
        rows.append([i, p["company"], p["product_name"], _cell(r, "body_type", idx, p["pid"], 40), _cell(r, "access_channel", idx, p["pid"], 30),
                     _cell(r, "launch_date", idx, p["pid"], 30), _cell(r, "channel", idx, p["pid"], 40), f"{ok}/{len(comp)}", "/".join(lv) or "—"])
    L += [caption("样本产品与资料概况"), "",
          _md_table(["序号", "公司", "产品", "适用人群定位", "健告类型", "上市时间", "销售渠道", "核心资料完整度", "证据等级"], rows), "",
          "注：核心资料完整度按条款、费率表、健康告知、责任免除/既往症清单、投保须知五类统计（“非缺文件”计为完整）；"
          "证据等级 A 为保险公司官网或备案原件，B 为持牌中介平台披露原件，C 为媒体报道或营销材料，D 为未核实线索。", ""]
    summary = (overall.get("summary") or "").strip()
    if summary:
        L += ["### （二）主要结论", ""] + say(summary, "overall.summary")

    # 二、承保范围
    L += ["## 二、承保范围", "", "### （一）核保端：健康告知与核保准入", ""]
    L += say("underwriting")
    L += [caption("健康告知与核保准入"), "", cmp_table(UNDERWRITING_KEYS, 160), "",
          "注：健告题数减少不等于准入放宽，需结合疾病问询范围与数据风控判断；免健告产品的投保限制可能体现在投保页或数据风控规则中。", ""]
    L += ["### （二）理赔端：既往症与赔付规则", ""]
    L += [caption("既往症处理与赔付规则"), "", cmp_table(CLAIM_KEYS, 200), ""]

    # 三、保障形态与版本演进
    L += ["## 三、保障形态与版本演进", "", "### （一）基础形态与费率", ""]
    L += say("overview") + say("renewal_pricing")
    L += [caption("基础形态、续保与费率"), "", cmp_table(FORM_KEYS, 120), "",
          "注：保费默认取 46–50 岁档、有社保、首年、年交、含等待期口径，具体口径见“保费口径与版本”。保险期间不等于保证续保期间。", ""]
    L += ["### （二）保障责任", ""]
    L += say("benefits")
    L += [caption("保障责任"), "", cmp_table(BENEFIT_KEYS, 120), ""]
    L += ["### （三）版本演进", ""]
    notes = analysis.get("version_notes") or {}
    if not cmp.get("versions"):
        L += ["本次样本未形成同系列新旧版本配对（或尚未运行 compare）。", ""]
    for i, v in enumerate(cmp.get("versions", []), 1):
        o, n = v["old"], v["new"]
        note = notes.get(f"{o['pid']}->{n['pid']}") or notes.get(f"{o['product_name']}->{n['product_name']}") or ""
        text = note if isinstance(note, str) else note.get("summary", "")
        if text:
            check_claims(text, f"version_notes {o['pid']}->{n['pid']}", problems)
        head = f"{o['product_name']}（{o.get('launch') or '上市时间待核'}）→{n['product_name']}（{n.get('launch') or '上市时间待核'}）。"
        L += [f"{i}．**{head}**{resolve_tokens(text, idx, problems) or '（研究者解读待撰写。）'}", ""]
        rows = []
        for ch in v["changes"]:
            orefs = fmt_refs(refs_for(idx, o["pid"], ch["field"]))
            nrefs = fmt_refs(refs_for(idx, n["pid"], ch["field"]))
            rows.append([ch["label"], truncate(ch["old"], 120) + orefs, truncate(ch["new"], 120) + nrefs, ch["direction"] or ch["change"]])
        if rows:
            L += [caption(f"{o['product_name']}与{n['product_name']}的主要变化"), "",
                  _md_table(["维度", "旧版", "新版", "变化（方向为候选判断）"], rows), ""]
        else:
            L += ["两版标准化字段未见差异。", ""]

    # 四、创新点识别
    L += ["## 四、创新点识别", ""]
    base = inn.get("baseline", {})
    L += [f"对每款产品的标准化字段做特征打标，与市场基线（样本摘要目录 {base.get('builtin_samples', 0)} 款、本次研究 "
          f"{base.get('workspace_products', 0)} 款）比较普及度、最早采用时间及是否为本系列新增。“首见”仅指样本内；"
          "仅有营销稿或媒体报道（C/D 级）支撑的特征不认定为创新；候选判定之外以研究者结论为准。", ""]
    judg = analysis.get("innovations") or {}
    k = 0
    for pid, info in (inn.get("products") or {}).items():
        jl = judg.get(pid) if isinstance(judg.get(pid), list) else []
        jl = [j for j in jl if (j.get("summary") or "").strip()]
        cands = [c for c in info["candidates"] if c["verdict"] != "行业常见/跟进" or c.get("new_in_series")]
        if not jl and not cands:
            continue
        L += [f"### （{CN_NUM[k] if k < 10 else k + 1}）{info['company']}·{info['product']}", ""]
        k += 1
        for i, j in enumerate(jl, 1):
            check_claims(j.get("summary", ""), f"innovations.{pid}.{j.get('feature')}", problems)
            body = " ".join(x for x in [j.get("summary", ""), ("风险逻辑：" + j["risk_logic"]) if j.get("risk_logic") else "",
                                        ("结论边界：" + j["boundary"]) if j.get("boundary") else ""] if x)
            refs = "".join(f"[{x}]" for x in j.get("refs", []))
            L += [f"{i}．**{j.get('label') or j.get('feature')}（{j.get('verdict', '')}）。**{resolve_tokens(body, idx, problems)}{refs}", ""]
        if cands:
            rows = []
            for c in cands:
                eref = ""
                for e in (c.get("evidence") or [])[:2]:
                    eref += fmt_refs(refs_for(idx, pid, e["field"])[:1])
                rows.append([c["label"], c["verdict"], {True: "是", False: "否", None: "—"}[c.get("new_in_series")],
                             f"{c['peers_with']}/{c['peers_total']}",
                             "；".join(f"{x['name']}（{x['launch'] or '—'}）" for x in c["peer_examples"][:2]) or "无",
                             truncate(c.get("product_snippet", ""), 60) + eref, "；".join(c.get("flags", [])) or "—"])
            L += [caption(f"{info['product']}的候选特征（{'脚本判定，供复核' if jl else '待研究者判断'}）"), "",
                  _md_table(["特征", "候选判定", "本系列新增", "样本普及度", "同业示例", "原文依据", "提示"], rows), ""]
    if k == 0:
        L += ["未发现样本内少见或本系列新增的特征。", ""]

    # 五、趋势判断与风险提示
    L += ["## 五、趋势判断与风险提示", ""]
    sub = 0
    finds = [f for f in analysis.get("findings") or [] if isinstance(f, dict) and (f.get("finding") or "").strip()]
    if finds:
        L += [f"### （{CN_NUM[sub]}）核心发现", ""]
        sub += 1
        for i, f in enumerate(finds, 1):
            body = " ".join(x for x in [f.get("detail", ""), ("风险控制逻辑与产品启示：" + f["logic"]) if f.get("logic") else "",
                                        ("结论边界：" + f["boundary"]) if f.get("boundary") else ""] if x)
            check_claims(f.get("detail", ""), f"findings[{i}]", problems)
            head = f["finding"].strip().rstrip("。") + "。"
            L += [f"{i}．**{head}**{resolve_tokens(body, idx, problems)}", ""]
    for key, lab in (("trends", "趋势判断"), ("risks", "风险提示")):
        lst = [t for t in overall.get(key) or [] if (t or "").strip()]
        if lst:
            L += [f"### （{CN_NUM[sub]}）{lab}", ""]
            sub += 1
            for i, t in enumerate(lst, 1):
                check_claims(t, f"overall.{key}", problems)
                L += [numbered(i, resolve_tokens(t, idx, problems)), ""]
    L += [f"### （{CN_NUM[sub]}）资料缺口与待核事项", ""]
    gaps = []
    for p in prods:
        r = recs[p["pid"]]
        pend = [fm[k]["label"] for k, f in r["fields"].items() if f.get("status") == "待核"]
        m = load_manifest(ws, p["pid"])
        miss = [f"{c} {DOC_TYPES[c]}：{s}" for c, s in completeness(m).items() if s != "已归档"]
        if pend or miss:
            gaps.append(f"{pname(p)}：" + "；".join(miss + ([f"待核字段：{'、'.join(pend)}"] if pend else [])) + "。")
    L += [numbered(i, g) for i, g in enumerate(gaps, 1)] if gaps else ["无。"]
    L.append("")
    if val:
        L += [f"校验摘要：error {val.get('totals', {}).get('errors', 0)} 项，warning {val.get('totals', {}).get('warnings', 0)} 项"
              "（明细见 outputs/校验报告.json 与对比表“校验结果”；error 未清零的内容不应作为结论引用）。", ""]

    # 附录
    L += ["## 附录A　证据清单", "", caption("证据清单", True), ""]
    rows = [[e["id"], f"{e['company']}·{e['product']}", e["label"], truncate(e.get("quote") or "", 120), e.get("doc"), e.get("page"),
             f"{e.get('match')}", e.get("evidence_level"), e.get("url") or "本地原件"] for e in items]
    L += [_md_table(["编号", "产品", "字段", "原文引文", "文件", "页码", "匹配", "证据等级", "来源"], rows), ""]
    L += ["## 附录B　资料清单与来源", "", caption("资料清单与来源", True), ""]
    rows = []
    for p in prods:
        for d in load_manifest(ws, p["pid"])["documents"]:
            src = (d.get("sources") or [{}])[0]
            rows.append([pname(p), d["doc_type_name"], d["file"], d.get("pages") or "", d.get("evidence_level"),
                         d["sha256"][:16], src.get("url") or "本地导入", (src.get("retrieved_at") or "")[:10]])
    L += [_md_table(["产品", "类型", "文件", "页数", "证据等级", "SHA256(前16位)", "来源链接", "获取日期"], rows), ""]
    if problems:
        L += ["## 附录C　报告自检提示", ""] + [f"{i}．{x}" for i, x in enumerate(problems, 1)] + [""]

    md = "\n".join(L)
    out_dir = out_dir or ws.outputs_dir
    stem = f"{title}_分析报告_{today().replace('-', '')}"
    md_path = Path(out_dir) / f"{stem}.md"
    write_text(md_path, md)
    result = {"markdown": str(md_path), "problems": problems}
    if docx:
        try:
            result["docx"] = str(markdown_to_docx(md, Path(out_dir) / f"{stem}.docx"))
        except ImportError:
            result["docx_skipped"] = "未安装 python-docx，仅输出 Markdown（pip install python-docx）"
    return result


# ---------------- Markdown → Word（研究报告体例）----------------

LATIN, EAST = "Times New Roman", "宋体"


def _set_fonts(rpr_owner, latin=LATIN, east=EAST):
    """同时设置西文与中文字体，并去掉主题字体属性（否则主题字体会覆盖显式字体）。"""
    from docx.oxml.ns import qn
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(a), None)
    for a, v in (("w:ascii", latin), ("w:hAnsi", latin), ("w:cs", latin), ("w:eastAsia", east)):
        rfonts.set(qn(a), v)


def _style(d, name, *, size=None, bold=None, align=None, indent=None, before=None, after=None, line=None,
           latin=LATIN, east=EAST, keep_next=False, base="Normal"):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt, RGBColor
    try:
        st = d.styles[name]
    except KeyError:
        st = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = d.styles[base]
        st.quick_style = True
    _set_fonts(st.element, latin, east)
    f = st.font
    f.color.rgb = RGBColor(0, 0, 0)
    f.italic = False
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
    pf.keep_with_next = keep_next or None
    return st


SETTINGS_AFTER_UPDATE = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                         "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                         "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                         "decimalSymbol", "listSeparator")


def _insert_before(parent, el, successors) -> None:
    """按 OOXML 的元素顺序插入（Word 对 tblPr / settings 子元素顺序较严格）。"""
    for child in parent:
        if child.tag.rsplit("}", 1)[-1] in successors:
            child.addprevious(el)
            return
    parent.append(el)


def _set_cell_border(cell, **kw):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, (sz, color) in kw.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single" if sz else "nil")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)


def _add_inline(par, text, size=None, bold_all=False):
    from docx.shared import Pt
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = par.add_run(part[2:-2] if bold else part)
        if bold or bold_all:
            run.bold = True
        if size:
            run.font.size = Pt(size)


def _field(par, instr: str, placeholder: str = ""):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def fld(t):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), t)
        return el
    r = par.add_run()
    r._r.append(fld("begin"))
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._r.append(it)
    r = par.add_run()
    r._r.append(fld("separate"))
    par.add_run(placeholder)
    r = par.add_run()
    r._r.append(fld("end"))


def _toc(d, heads: list[tuple[int, str]]) -> None:
    """目录域：预先写入各级标题（打开文档时 Word 提示更新域即可补全页码）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def fld(par, t):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), t)
        par.add_run()._r.append(el)
    heads = heads or [(1, "目录")]
    first = d.add_paragraph(style=f"toc {heads[0][0]}")
    fld(first, "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    first.add_run()._r.append(it)
    fld(first, "separate")
    first.add_run(heads[0][1])
    last = first
    for lvl, t in heads[1:]:
        last = d.add_paragraph(style=f"toc {lvl}")
        last.add_run(t)
    fld(last, "end")


def _col_widths(header: list[str], body: list[list[str]], total_cm: float = 16.0) -> list[float]:
    """按各列平均字数分配列宽；短列（序号、页码）给窄列，标签列与产品名列保证最小宽度。"""
    weights, mins = [], []
    for c in range(len(header)):
        lens = [len(r[c]) if c < len(r) else 0 for r in body[:60]] or [0]
        avg = sum(lens) / len(lens)
        weights.append(min(max(avg, 3.0), 45.0))
        mins.append(0.9 if max(lens + [len(header[c]) / 2]) <= 4 else 2.3 if c == 0 or header[c] in ("产品", "公司") else 1.5)
    s = sum(weights)
    w = [max(m, total_cm * x / s) for x, m in zip(weights, mins)]
    k = total_cm / sum(w)
    return [x * k for x in w]


def markdown_to_docx(md: str, path: Path) -> Path:
    import docx  # type: ignore
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)
    J, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER
    _style(d, "Normal", size=10.5, line=18)
    _style(d, "Report Title CN", size=20, bold=True, align=C, after=6, latin="微软雅黑", east="微软雅黑")
    _style(d, "TOC Title CN", size=14, bold=True, align=C, after=12, line=21)
    _style(d, "Heading 1", size=14, bold=True, before=12, after=6, line=21, keep_next=True)
    _style(d, "Heading 2", size=11, bold=True, before=6, after=3, line=19, keep_next=True)
    _style(d, "Body Text CN", align=J, indent=21, after=2)
    _style(d, "Table Caption CN", size=10, bold=True, align=C, before=8, after=4, line=15, keep_next=True)
    _style(d, "Note CN", size=8.5, align=J, before=3, after=8, line=12.5)
    _style(d, "Table Text CN", size=7.5, line=11.5)
    _style(d, "toc 1", size=11, bold=True, before=6, line=20)
    _style(d, "toc 2", size=10.5, line=20).paragraph_format.left_indent = Pt(21)
    # 页脚页码
    fp = sec.footer.paragraphs[0]
    fp.alignment = C
    _field(fp, "PAGE", "1")
    for r in fp.runs:
        r.font.size = Pt(9)
    # 打开文档时更新目录域
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    _insert_before(d.settings.element, upd, SETTINGS_AFTER_UPDATE)

    lines = md.splitlines()
    i, toc_done = 0, False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|(?:-+\|)+$", lines[i + 1]):
            header = [c.strip() for c in ln.strip().strip("|").split("|")]
            body = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                body.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = d.add_table(rows=1 + len(body), cols=len(header))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            tblPr = t._tbl.tblPr
            mar = OxmlElement("w:tblCellMar")
            for edge, w in (("top", 30), ("left", 60), ("bottom", 30), ("right", 60)):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), str(w))
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            _insert_before(tblPr, mar, ("tblLook", "tblCaption", "tblDescription"))
            widths = _col_widths(header, body)
            label_col = header[0] in ("维度", "产品", "项目", "序号", "编号")
            for r_i, row in enumerate([header] + body):
                if r_i == 0:
                    trPr = t.rows[0]._tr.get_or_add_trPr()
                    th = OxmlElement("w:tblHeader")
                    th.set(qn("w:val"), "true")
                    trPr.append(th)
                for c_i in range(len(header)):
                    cell = t.cell(r_i, c_i)
                    cell.width = Cm(widths[c_i])
                    txt = (row[c_i] if c_i < len(row) else "").replace("<br>", "\n")
                    par = cell.paragraphs[0]
                    par.style = d.styles["Table Text CN"]
                    head = r_i == 0
                    top = (12, "000000") if head else (0, "FFFFFF")
                    bottom = (6, "000000") if head else ((12, "000000") if r_i == len(body) else (2, "BFBFBF"))
                    _set_cell_border(cell, top=top, bottom=bottom, left=(0, "FFFFFF"), right=(0, "FFFFFF"))
                    if head or (label_col and c_i == 0):
                        par.alignment = C
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _add_inline(par, txt, 7.5, bold_all=head or (label_col and c_i == 0 and header[0] != "编号"))
            continue
        if ln.startswith("# "):
            _add_inline(d.add_paragraph(style="Report Title CN"), ln[2:])
        elif ln.startswith("## "):
            if not toc_done:
                toc_done = True
                d.add_paragraph("目　录", style="TOC Title CN")
                _toc(d, [(1 if x.startswith("## ") else 2, x.split(" ", 1)[1]) for x in lines if x.startswith(("## ", "### "))])
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            _add_inline(d.add_paragraph(style="Heading 1"), ln[3:])
        elif ln.startswith("### "):
            _add_inline(d.add_paragraph(style="Heading 2"), ln[4:])
        elif ln.startswith("> "):
            p = d.add_paragraph(style="Note CN")
            p.alignment = C
            _add_inline(p, ln[2:])
        elif CAPTION.match(ln):
            _add_inline(d.add_paragraph(style="Table Caption CN"), ln)
        elif ln.startswith("注："):
            _add_inline(d.add_paragraph(style="Note CN"), ln)
        elif ln.strip():
            _add_inline(d.add_paragraph(style="Body Text CN"), ln[2:] if ln.startswith("- ") else ln)
        i += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        d.save(str(path))
    except PermissionError:
        path = path.with_name(path.stem + "_new" + path.suffix)
        d.save(str(path))
    return path
