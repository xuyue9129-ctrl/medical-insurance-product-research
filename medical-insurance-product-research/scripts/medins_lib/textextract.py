"""把 PDF / HTML / DOCX / XLSX / TXT 转成“按页”的文本，保留页码用于引文定位。

PDF 依次尝试 PyMuPDF(fitz) → pdfplumber → pypdf；都不可用时报错并提示安装。
扫描件（每页字符极少）会被标记 scanned=True，提示需要 OCR 或人工核对，不能当作已核验文本。
"""
from __future__ import annotations

import html as _html
import re
from pathlib import Path

from .common import MedinsError, read_text

PAGE_MARK = "=== 第{n}页 ==="
PAGE_MARK_RE = re.compile(r"^=== 第(\d+)页 ===$", re.M)


def available_pdf_backends() -> list[str]:
    out = []
    for mod, name in (("fitz", "pymupdf"), ("pdfplumber", "pdfplumber"), ("pypdf", "pypdf")):
        try:
            __import__(mod)
            out.append(name)
        except Exception:
            pass
    return out


def pdf_page_count(path: Path) -> int | None:
    try:
        import fitz  # type: ignore
        with fitz.open(str(path)) as doc:
            return doc.page_count
    except Exception:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        return len(PdfReader(str(path)).pages)
    except Exception:
        pass
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            return len(pdf.pages)
    except Exception:
        return None


def _fitz_rows(page) -> str:
    """按行重建（同一水平线上的单元格拼成一行），用于费率表等表格型 PDF。"""
    words = page.get_text("words")
    rows: list[list] = []
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        yc, h = (w[1] + w[3]) / 2, (w[3] - w[1])
        if rows and abs(rows[-1][0] - yc) <= max(2.0, h * 0.45):
            rows[-1][1].append(w)
        else:
            rows.append([yc, [w]])
    return "\n".join(" ".join(x[4] for x in sorted(r[1], key=lambda w: w[0])) for r in rows)


def _pdf_pages(path: Path, layout: bool = False) -> tuple[list[str], str]:
    errors = []
    try:
        import fitz  # type: ignore
        with fitz.open(str(path)) as doc:
            if layout:
                return [_fitz_rows(p) for p in doc], "pymupdf-rows"
            return [p.get_text("text") for p in doc], "pymupdf"
    except ImportError:
        pass
    except Exception as e:  # 损坏/加密
        errors.append(f"pymupdf: {e}")
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            return [(p.extract_text() or "") for p in pdf.pages], "pdfplumber"
    except ImportError:
        pass
    except Exception as e:
        errors.append(f"pdfplumber: {e}")
    try:
        from pypdf import PdfReader  # type: ignore

        def _pp(p):
            if layout:
                try:
                    return p.extract_text(extraction_mode="layout") or ""
                except TypeError:  # 旧版 pypdf 不支持 layout 模式
                    pass
            return p.extract_text() or ""
        return [_pp(p) for p in PdfReader(str(path)).pages], "pypdf"
    except ImportError:
        pass
    except Exception as e:
        errors.append(f"pypdf: {e}")
    if errors:
        raise MedinsError(f"PDF 无法读取（可能损坏或加密）：{path.name}；" + "；".join(errors))
    raise MedinsError("未安装 PDF 解析库：请 `pip install pymupdf` 或 `pip install pdfplumber` 或 `pip install pypdf`")


def pdf_tables(path: Path, max_pages: int = 40) -> list[dict]:
    """尽力抽取表格（费率表用）；需要 pdfplumber，缺失时返回空列表。"""
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []
    out = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages[:max_pages], 1):
                for t in page.extract_tables() or []:
                    rows = [[(c or "").strip() for c in row] for row in t if any(c for c in row)]
                    if rows:
                        out.append({"page": i, "rows": rows})
    except Exception:
        return out
    return out


def _html_text(raw: str) -> str:
    try:
        from bs4 import BeautifulSoup  # type: ignore
        soup = BeautifulSoup(raw, "html.parser")
        for t in soup(["script", "style", "noscript", "svg"]):
            t.decompose()
        text = soup.get_text("\n")
    except ImportError:
        text = re.sub(r"(?is)<(script|style|noscript).*?</\1>", " ", raw)
        text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</tr>|</h\d>", "\n", text)
        text = _html.unescape(re.sub(r"<[^>]+>", " ", text))
    text = re.sub(r"[ \t\u3000]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n", text).strip()


def _docx_text(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as e:
        raise MedinsError("读取 .docx 需要 `pip install python-docx`") from e
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx_pages(path: Path) -> list[str]:
    import openpyxl  # type: ignore
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    pages = []
    for ws in wb.worksheets:
        lines = [f"[工作表] {ws.title}"]
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v not in (None, "")]
            if vals:
                lines.append(" | ".join(vals))
        pages.append("\n".join(lines))
    return pages


def extract_pages(path: Path, layout: bool = False) -> dict:
    """返回 {pages: [str], backend, scanned, chars}。layout=True 时按行重建（费率表）。"""
    path = Path(path)
    ext = path.suffix.lower()
    backend = "text"
    if ext == ".pdf":
        pages, backend = _pdf_pages(path, layout)
    elif ext in (".html", ".htm", ".mhtml"):
        pages, backend = [_html_text(read_text(path))], "html"
    elif ext == ".docx":
        pages, backend = [_docx_text(path)], "python-docx"
    elif ext in (".xlsx", ".xlsm"):
        pages, backend = _xlsx_pages(path), "openpyxl"
    elif ext in (".txt", ".md", ".json", ".csv"):
        pages = [read_text(path)]
    else:
        raise MedinsError(f"暂不支持的文件类型：{path.name}（支持 pdf/html/docx/xlsx/txt）")
    pages = [p.replace("\x00", "") for p in pages]
    chars = sum(len(re.sub(r"\s", "", p)) for p in pages)
    scanned = ext == ".pdf" and len(pages) > 0 and chars < 40 * len(pages)
    return {"pages": pages, "backend": backend, "scanned": scanned, "chars": chars}


def pages_to_text(pages: list[str]) -> str:
    return "\n".join(f"{PAGE_MARK.format(n=i)}\n{p.strip()}\n" for i, p in enumerate(pages, 1))


def text_to_pages(text: str) -> list[str]:
    """读取带页码标记的缓存文本；没有标记时整体视为第1页。"""
    marks = list(PAGE_MARK_RE.finditer(text))
    if not marks:
        return [text]
    pages: dict[int, str] = {}
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        pages[int(m.group(1))] = text[m.end():end].strip("\n")
    n = max(pages)
    return [pages.get(i, "") for i in range(1, n + 1)]
