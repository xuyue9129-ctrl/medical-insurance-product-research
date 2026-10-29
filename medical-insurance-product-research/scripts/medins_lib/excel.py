"""生成 Excel 对比表。按样本自动选择两种版式（--mode 可强制指定）：

series（同一系列的多个版本/计划）：
  说明 | ① 系列总览 | ② 保障责任明细 | ③ 保费对比 | ④ 健告与既往症 | ⑤ 版本对比 | ⑥ 创新点识别 | ⑦ 证据索引 | ⑧ 资料清单 | ⑨ 校验结果
market（不同公司/系列的多产品横比）：
  说明 | ① 产品总览 | ② 严重既往症对比 | ③ 产品演进与趋势 | ④ 核心发现 | ⑤ 维度对比 | ⑥ 保障责任明细 | ⑦ 保费对比 |
  ⑧ 健告与既往症 | ⑨ 创新点识别 | ⑩ 证据索引 | ⑪ 资料清单 | ⑫ 校验结果
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from .citations import build_index, fmt_refs, refs_for
from .common import MedinsError, read_json, today, truncate
from .doctypes import DOC_TYPES
from .fetch import EVIDENCE_LEVELS, completeness, load_manifest
from .records import field_map, load_record, schema_fields
from .workspace import Workspace

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
except ImportError as e:  # pragma: no cover
    raise MedinsError("生成 Excel 需要 openpyxl：pip install openpyxl") from e

FONT = "微软雅黑"
NAVY, TEAL, SECTION, LABEL, KEY, GREY_TXT = "1F3864", "1F6F8B", "D9E1F2", "F2F2F2", "DCEBEE", "7F7F7F"
BAR, HEAD, GROUP, INK, BAND = "173B5D", "245780", "DDEBF7", "263440", "F3F6F9"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_STYLE = {"推定": ("1F6F8B", True), "待核": (GREY_TXT, False), "不适用": (GREY_TXT, False),
                "待复核": ("C00000", False), "规则抽取": ("7F6000", False)}
MARKS = {"✔": ("C6EFCE", "006100"), "✘": ("FFC7CE", "9C0006"), "◐": ("FFEB9C", "9C5700")}
MAX_CELL = 32000

# 多产品总览的列分组（字段顺序即列顺序）
MARKET_GROUPS = [
    ("产品基本信息", ["company", "product_name", "body_type", "launch_date", "channel", "term", "guaranteed_renewal"]),
    ("投保与健康告知", ["access_channel", "hd_count", "entry_age", "nodules_three_highs"]),
    ("既往症与理赔", ["preexisting_special", "first_year_rules", "waiting_period", "deductible", "coinsurance", "no_social_insurance"]),
    ("核心额度与特殊责任", ["total_limit", "rehab", "outpatient", "external_drugs", "advanced_therapy"]),
    ("保费", ["premium_50m", "premium_basis"]),
]
NARROW = {"company", "launch_date", "term", "guaranteed_renewal", "hd_count"}
MEDIUM = {"access_channel", "body_type", "entry_age", "channel", "waiting_period"}
# 责任覆盖速查（✔ 含 / ✘ 不含 / ◐ 待核）
COVERAGE_KEYS = ["general_medical", "critical_medical", "special_outpatient", "outpatient", "rehab", "external_drugs",
                 "advanced_therapy", "health_services"]
# 严重既往症病种速查（关键词命中清单原文；以原文为准）
DISEASE_GROUPS = [
    ("恶性肿瘤", r"恶性肿瘤|癌|白血病|淋巴瘤|肉瘤"),
    ("肝占位/肝硬化", r"肝[^，。；]{0,6}占位|肝硬化|肝功能(?:衰竭|失代偿)"),
    ("结节/肿块", r"结节|肿块|息肉|囊肿"),
    ("肾脏疾病", r"肾|尿毒症|透析"),
    ("心脑血管", r"冠心病|心肌梗|心绞痛|心力衰竭|心功能|脑梗|脑出血|脑卒中|中风|动脉瘤"),
    ("三高及并发症", r"高血压|糖尿病|高脂血"),
    ("自身免疫", r"红斑狼疮|类风湿|自身免疫|强直性脊柱炎|硬皮病"),
    ("其他重大疾病", r"移植|帕金森|阿尔茨海默|慢性阻塞性肺|肺心病|瘫痪|精神分裂|再生障碍"),
]
# 版本演进按维度归并
EVOLUTION_COLS = [("健告与准入变化", {"投保与核保准入"}), ("既往症与赔付变化", {"既往症与赔付规则"}),
                  ("保障责任变化", {"核心额度与保障责任", "特殊责任与健康服务"}), ("续保、保费及其他变化", {"基本信息", "续保机制", "保费", "定位与资料"})]


def _font(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _put(ws, r, c, v, bold=False, fill=None, color="000000", size=10, center=False, italic=False, wrap=True, link=None):
    v = truncate(v, MAX_CELL) if isinstance(v, str) else v
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = _font(bold, color, size, italic)
    cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center" if center else "top", wrap_text=wrap)
    cell.border = BORDER
    if fill:
        cell.fill = _fill(fill)
    if link:
        cell.hyperlink = link
        cell.font = Font(name=FONT, size=size, color="0563C1", underline="single")
    return cell


def _title(ws, text, sub="", span=8):
    ws.cell(row=1, column=1, value=text).font = _font(True, NAVY, 16)
    if sub:
        c = ws.cell(row=2, column=1, value=sub)
        c.font = _font(False, "595959", 9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, span))
        ws.row_dimensions[2].height = 30


def _bar(ws, text, span, note=""):
    """总表式标题栏：第 1 行深色底白字，跨全部列。"""
    for c in range(1, span + 1):
        ws.cell(row=1, column=c).fill = _fill(BAR)
    cell = ws.cell(row=1, column=1, value=text + (f"　｜　{note}" if note else ""))
    cell.font = _font(True, "FFFFFF", 14)
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, span))
    ws.row_dimensions[1].height = 30


def _groups(ws, r, start_col, groups, fill=GROUP, color=BAR):
    col = start_col
    for name, n in groups:
        for k in range(n):
            _put(ws, r, col + k, name if k == 0 else "", bold=True, fill=fill, color=color, center=True)
        if n > 1:
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + n - 1)
        col += n
    ws.row_dimensions[r].height = 20


def _header(ws, r, headers, fill=NAVY):
    for i, h in enumerate(headers, 1):
        _put(ws, r, i, h, bold=True, fill=fill, color="FFFFFF", center=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _autoheight(ws, first_row, last_row, min_h=18, max_h=409):
    widths = {}
    for col, dim in ws.column_dimensions.items():
        widths[col] = dim.width or 10
    for r in range(first_row, last_row + 1):
        lines = 1
        for cell in ws[r]:
            if cell.value is None or not isinstance(cell.value, str):
                continue
            w = widths.get(cell.column_letter, 10)
            per_line = max(4, int(w * 0.9 / 1.9) * 1.0)  # 中文约占 2 个字符宽
            n = sum(max(1, math.ceil(len(seg) / per_line)) for seg in cell.value.split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = max(min_h, min(max_h, 14.5 * lines + 4))


def _cell_text(f: dict, refs: list[str] | None = None) -> str:
    st = f.get("status", "未填")
    v = f.get("value", "")
    if st == "未填":
        return ""
    if st in ("待核", "不适用"):
        return f"〔{st}〕{v + '；' if v else ''}{f.get('note', '')}"
    tag = {"推定": "〔推定〕", "待复核": "〔待复核〕", "规则抽取": "〔规则抽取·待确认〕"}.get(st, "")
    return f"{v}{(' ' + tag) if tag else ''}{(' ' + fmt_refs(refs)) if refs else ''}"


def _style_status(cell, st, size=10, color=None):
    if st in STATUS_STYLE:
        c, italic = STATUS_STYLE[st]
        cell.font = _font(False, c, size, italic)
    elif color:
        cell.font = _font(False, color, size)


def _mark_cell(cell, mark):
    key = mark[:1]
    if key in MARKS:
        bg, fg = MARKS[key]
        cell.fill = _fill(bg)
        cell.font = _font(True, fg)


def coverage_mark(fld: dict) -> str:
    """责任覆盖速查：✔ 含 / ✘ 不含 / ◐ 待核 / — 未填。"""
    st, v = fld.get("status", "未填"), (fld.get("value") or "").strip()
    if st == "未填":
        return "—"
    if st == "待核":
        return "◐ 待核"
    if st == "不适用" or re.match(r"^(无|不含|不承担|不保障|未包含|未提供|不提供)", v):
        return "✘"
    return "✔"


def disease_marks(rec: dict) -> list[str]:
    """严重既往症病种速查：按关键词判断清单原文是否涉及该类疾病。"""
    srcs = [rec["fields"].get(k, {}) for k in ("preexisting_serious", "preexisting_special")]
    text = "".join(f.get("value", "") for f in srcs if f.get("status") not in ("待核", "不适用", "未填"))
    if not text:
        st = srcs[0].get("status", "未填")
        return ["◐ 待核" if st == "待核" else ("✘ 不除外" if st == "不适用" else "—")] * len(DISEASE_GROUPS)
    return ["✔ 涉及" if re.search(pat, text) else "—" for _, pat in DISEASE_GROUPS]


def _series_key(p: dict) -> tuple:
    return (p.get("company"), p.get("series_line") or p.get("series") or p.get("product_name"))


def detect_mode(prods: list[dict]) -> str:
    return "series" if len({_series_key(p) for p in prods}) == 1 else "market"


def _finish(ws, landscape=True):
    ws.sheet_view.zoomScale = 90
    wide = ws.max_column > 14
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if wide else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = math.ceil(ws.max_column / 14) if wide else 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True


def build_workbook(ws_: Workspace, pids: list[str] | None = None, out: Path | None = None, title: str = "",
                   mode: str = "auto") -> Path:
    from .report import resolve_tokens
    pids = [ws_.get(p)["pid"] for p in (pids or list(ws_.products()))]
    if not pids:
        raise MedinsError("工作区没有产品")
    prods = [ws_.get(p) for p in pids]
    recs = [load_record(ws_, p) for p in pids]
    fm = field_map()
    ev_items, ev_idx = build_index(ws_, pids)
    cmp = read_json(ws_.outputs_dir / "对比结果.json", {}) or {}
    inn = read_json(ws_.outputs_dir / "创新点候选.json", {}) or {}
    val = read_json(ws_.outputs_dir / "校验报告.json", {}) or {}
    analysis = read_json(ws_.analysis_path, {}) or {}
    if not isinstance(analysis, dict):
        analysis = {}
    title = title or (ws_.load() or {}).get("title") or "医疗险产品研究"
    mode = detect_mode(prods) if mode in (None, "", "auto") else mode
    if mode not in ("series", "market"):
        raise MedinsError("mode 只能是 auto / series / market")
    names = [f"{p['company']}\n{p['product_name']}" for p in prods]
    sections = (read_json(Path(__file__).resolve().parents[2] / "assets" / "schema" / "fields.json", {}) or {}).get("sections", [])

    def text(s: str) -> str:
        return resolve_tokens(s or "", ev_idx, [])

    def fcell(s, r, c, k, key, **kw):
        fld = recs[k]["fields"][key]
        cell = _put(s, r, c, _cell_text(fld, refs_for(ev_idx, prods[k]["pid"], key)), **kw)
        _style_status(cell, fld.get("status"), kw.get("size", 10), kw.get("color"))
        return cell

    wb = Workbook()
    wb.remove(wb.active)
    order = (["说明", "系列总览", "保障责任明细", "保费对比", "健告与既往症", "版本对比", "创新点识别", "证据索引", "资料清单", "校验结果"]
             if mode == "series" else
             ["说明", "产品总览", "严重既往症对比", "产品演进与趋势", "核心发现", "维度对比", "保障责任明细", "保费对比", "健告与既往症",
              "创新点识别", "证据索引", "资料清单", "校验结果"])
    circled = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮"
    sheet_name = {n: (n if i == 0 else f"{circled[i - 1]} {n}") for i, n in enumerate(order)}

    # ---------- 说明 ----------
    def sh_readme():
        s = wb.create_sheet(sheet_name["说明"])
        kind = "同系列版本对比" if mode == "series" else "多产品横向对比"
        _title(s, f"{title} · 对比表", f"生成日期 {today()}；{kind}，共 {len(prods)} 款产品。所有取值均来自已归档原件，单元格末尾 [E#] 为证据编号，"
                                       f"对应“{sheet_name['证据索引']}”。")
        rows = [("工作表", "内容")] + [(sheet_name[n], SHEET_DESC[n]) for n in order[1:]] + [
                ("", ""), ("字段状态", "含义"),
                ("（无标记）", "原文核验：引文已在归档原件中逐字/高相似度找到"),
                ("〔规则抽取·待确认〕", "脚本按规则抽出，引文可核验但含义需研究者确认"),
                ("〔推定〕（蓝色斜体）", "基于原文的归纳或推断，依据见证据索引/备注"),
                ("〔待核〕（灰色）", "未取得资料或无法确认，后附原因"),
                ("〔待复核〕（红色）", "引文暂未在原文中找到，不得作为结论"),
                ("✔ / ✘ / ◐", "速查标记：含（涉及）/ 不含 / 待核；速查只作导航，以原文与证据为准"),
                ("浅青底色维度", "该维度各产品取值存在差异"),
                ("", ""), ("证据等级", "说明")] + [(k, v) for k, v in EVIDENCE_LEVELS.items()] + [
                ("", ""), ("口径提醒", "保险期间≠保证续保；健告题数减少≠准入放宽；上市时间以发布会/官网上线稿为准，附件配置时间不等于首发日；"
                                      "保费默认 46–50 岁档、有社保、首年、年交、含等待期，组合产品分列主险与附加险。")]
        for i, (a, b) in enumerate(rows, 4):
            h = i == 4 or a in ("字段状态", "证据等级")
            _put(s, i, 1, a, bold=bool(a), fill=NAVY if h else (LABEL if a else None), color="FFFFFF" if h else "000000")
            _put(s, i, 2, b, bold=h, fill=NAVY if h else None, color="FFFFFF" if h else "000000")
        _widths(s, [24, 110])
        _autoheight(s, 4, 4 + len(rows))
        _finish(s, False)

    # ---------- 系列总览（维度 × 产品）----------
    def sh_series_overview():
        s = wb.create_sheet(sheet_name["系列总览"])
        _title(s, f"{title} · 系列总览", "产品为列、维度为行；浅青底色表示各版本取值不同。“责任覆盖速查”与“保费速览”便于快速定位差异，细节见后续工作表。",
               len(prods) + 1)
        _header(s, 4, ["对比维度"] + names)
        r = 5

        def section(name):
            nonlocal r
            _put(s, r, 1, name, bold=True, fill=SECTION, color=NAVY)
            for j in range(len(prods)):
                _put(s, r, j + 2, "", fill=SECTION)
            s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(prods) + 1)
            r += 1

        for sec in sections:
            keys = [x["key"] for x in schema_fields() if x["section"] == sec]
            if sec == "保费":
                section("保费速览")
            else:
                section(sec)
            for key in keys:
                vals = [rec["fields"][key] for rec in recs]
                filled = {re.sub(r"\s+", "", v.get("value", "")) for v in vals if v.get("value")}
                _put(s, r, 1, fm[key]["label"], bold=True, fill=KEY if len(filled) > 1 else LABEL)
                for k in range(len(prods)):
                    fcell(s, r, k + 2, k, key, fill="E3F1F5" if len(filled) > 1 else None)
                r += 1
            if sec == "核心额度与保障责任":
                section("责任覆盖速查（✔ 含 / ✘ 不含 / ◐ 待核）")
                for key in COVERAGE_KEYS:
                    _put(s, r, 1, fm[key]["label"], bold=True, fill=LABEL)
                    for k, rec in enumerate(recs):
                        mk = coverage_mark(rec["fields"][key])
                        _mark_cell(_put(s, r, k + 2, mk, center=True), mk)
                    r += 1
                bnames = []
                for rec in recs:
                    for row in rec.get("benefits") or []:
                        if row["name"] not in bnames:
                            bnames.append(row["name"])
                for bn in bnames:
                    _put(s, r, 1, f"责任明细·{bn}", bold=True, fill=LABEL)
                    for k, rec in enumerate(recs):
                        has = any(x["name"] == bn for x in rec.get("benefits") or [])
                        mk = "✔" if has else ("—" if not rec.get("benefits") else "✘")
                        _mark_cell(_put(s, r, k + 2, mk, center=True), mk)
                    r += 1
        _widths(s, [24] + [48 if len(prods) <= 3 else 38] * len(prods))
        s.freeze_panes = "B5"
        _autoheight(s, 5, r - 1)
        _finish(s)

    # ---------- 产品总览（每行一款，总表式）----------
    def sh_market_overview():
        s = wb.create_sheet(sheet_name["产品总览"])
        cols = [k for _, ks in MARKET_GROUPS for k in ks]
        span = len(cols) + 3
        _bar(s, f"{title} · 产品总览", span, f"{len(prods)} 款产品；[E#] 为证据编号")
        groups = [("", 1)] + [(g, len(ks)) for g, ks in MARKET_GROUPS] + [("资料与证据", 2)]
        _groups(s, 2, 1, groups)
        _header(s, 3, ["序号"] + [fm[k]["label"] for k in cols] + ["核心资料完整度", "证据等级分布"], fill=HEAD)
        for i, (p, rec) in enumerate(zip(prods, recs), 1):
            r = 3 + i
            band = BAND if i % 2 == 0 else None
            _put(s, r, 1, i, center=True, fill=band, color=INK)
            for j, key in enumerate(cols, 2):
                fcell(s, r, j, i - 1, key, fill=band, color=INK, bold=key in ("company", "product_name"),
                      center=key in ("company", "launch_date", "hd_count"))
            m = load_manifest(ws_, p["pid"])
            comp = completeness(m)
            _put(s, r, len(cols) + 2, "\n".join(f"{k} {DOC_TYPES[k].split('/')[0]}：{v}" for k, v in comp.items()), fill=band, color=INK, size=9)
            lv = {}
            for d in m["documents"]:
                lv[d.get("evidence_level", "?")] = lv.get(d.get("evidence_level", "?"), 0) + 1
            _put(s, r, len(cols) + 3, "；".join(f"{k}级{v}份" for k, v in sorted(lv.items())) or "无", fill=band, color=INK, size=9)
        _widths(s, [6] + [(13 if k in NARROW else 22 if k in MEDIUM else 34) for k in cols] + [24, 14])
        s.column_dimensions["C"].width = 24
        s.freeze_panes = "D4"
        s.auto_filter.ref = f"A3:{get_column_letter(span)}{3 + len(prods)}"
        _autoheight(s, 4, 3 + len(prods))
        _finish(s)

    # ---------- 严重既往症对比 ----------
    def sh_preexisting():
        s = wb.create_sheet(sheet_name["严重既往症对比"])
        keys = ["access_channel", "preexisting_definition", "preexisting_general", "preexisting_serious", "preexisting_special", "first_year_rules"]
        span = 3 + len(keys) + len(DISEASE_GROUPS)
        _bar(s, "严重既往症与既往症处理对比", span, "病种速查按清单原文关键词标注，仅作导航，以原文为准")
        _groups(s, 2, 1, [("", 3), ("既往症处理（原文要点）", len(keys)), ("严重既往症病种速查", len(DISEASE_GROUPS))])
        _header(s, 3, ["序号", "保险公司", "产品"] + [fm[k]["label"] for k in keys] + [g for g, _ in DISEASE_GROUPS], fill=HEAD)
        for i, (p, rec) in enumerate(zip(prods, recs), 1):
            r = 3 + i
            band = BAND if i % 2 == 0 else None
            _put(s, r, 1, i, center=True, fill=band, color=INK)
            _put(s, r, 2, p["company"], bold=True, center=True, fill=band, color=INK)
            _put(s, r, 3, p["product_name"], bold=True, fill=band, color=INK)
            for j, key in enumerate(keys, 4):
                fcell(s, r, j, i - 1, key, fill=band, color=INK)
            for j, mk in enumerate(disease_marks(rec), 4 + len(keys)):
                _mark_cell(_put(s, r, j, mk, center=True, fill=band, color=INK), mk)
        _widths(s, [6, 12, 24, 12] + [40] * (len(keys) - 1) + [11] * len(DISEASE_GROUPS))
        s.freeze_panes = "D4"
        _autoheight(s, 4, 3 + len(prods))
        _finish(s)

    # ---------- 产品演进与趋势 ----------
    def sh_evolution():
        s = wb.create_sheet(sheet_name["产品演进与趋势"])
        heads = ["产品系列", "保险公司", "版本演进"] + [c for c, _ in EVOLUTION_COLS] + ["研究者解读/趋势"]
        _bar(s, "同系列产品版本变化", len(heads), "变化要点由标准化字段比对生成；方向为候选判断，需结合条款复核")
        _header(s, 2, heads, fill=HEAD)
        notes = analysis.get("version_notes") or {}
        r = 3
        for v in cmp.get("versions", []):
            o, n = v["old"], v["new"]
            p = ws_.get(n["pid"])
            _put(s, r, 1, p.get("series_line") or p.get("series") or n["product_name"], bold=True, color=HEAD)
            _put(s, r, 2, p["company"], center=True, color=INK)
            _put(s, r, 3, f"{o['product_name']}（{o.get('launch') or '上市时间待核'}）\n→ {n['product_name']}（{n.get('launch') or '上市时间待核'}）", color=INK)
            for j, (_, secs) in enumerate(EVOLUTION_COLS, 4):
                lines = []
                for ch in v["changes"]:
                    if ch.get("section") not in secs:
                        continue
                    if ch["change"] == "新增":
                        line = f"· {ch['label']}（新增）：{truncate(ch['new'], 40)}"
                    elif ch["change"] != "变化":
                        line = f"· {ch['label']}（{ch['change']}）：原为{truncate(ch['old'], 30)}"
                    else:
                        line = f"· {ch['label']}：{truncate(ch['old'], 28)} → {truncate(ch['new'], 28)}"
                    lines.append(line + (f"【{ch['direction']}】" if ch.get("direction") else ""))
                more = f"\n……另 {len(lines) - 8} 项见“版本对比”" if len(lines) > 8 else ""
                _put(s, r, j, ("\n".join(lines[:8]) + more) or "—", color=INK)
            note = notes.get(f"{o['pid']}->{n['pid']}") or notes.get(f"{o['product_name']}->{n['product_name']}") or ""
            _put(s, r, len(heads), text(note if isinstance(note, str) else note.get("summary", "")) or "（待研究者填写 analysis.json → version_notes）",
                 color=INK if note else GREY_TXT)
            r += 1
        if r == 3:
            _put(s, 3, 1, "样本中没有同系列新旧版本配对（可用 compare --pair 旧pid:新pid 指定）。", color=GREY_TXT)
            s.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(heads))
        _widths(s, [20, 12, 30, 36, 36, 36, 36, 44])
        s.freeze_panes = "C3"
        _autoheight(s, 3, max(3, r - 1))
        _finish(s)

    # ---------- 核心发现 ----------
    def sh_findings():
        s = wb.create_sheet(sheet_name["核心发现"])
        heads = ["核心发现", "具体产品、阶段与变化", "风险控制逻辑与产品启示", "结论边界"]
        _bar(s, "核心发现（研究者结论）", len(heads), "来自 analysis.json 的 findings 与已判断的创新点；[E#] 为证据编号")
        _header(s, 2, heads, fill=HEAD)
        rows = []
        for f in analysis.get("findings") or []:
            if isinstance(f, dict) and (f.get("finding") or "").strip():
                rows.append([f.get("finding", ""), text(f.get("detail", "")), text(f.get("logic", "")), f.get("boundary", "")])
        prod_name = {p["pid"]: f"{p['company']}·{p['product_name']}" for p in prods}
        for pid, lst in (analysis.get("innovations") or {}).items():
            for j in lst if isinstance(lst, list) else []:
                if (j.get("summary") or "").strip() and str(j.get("verdict", "")).startswith(("真创新", "差异化")):
                    rows.append([f"{j.get('label') or j.get('feature')}（{j['verdict']}）", f"{prod_name.get(pid, pid)}：{text(j['summary'])}",
                                 text(j.get("risk_logic", "")), j.get("boundary", "")])
        for i, row in enumerate(rows, 3):
            band = BAND if i % 2 == 0 else None
            for k, x in enumerate(row, 1):
                _put(s, i, k, x, bold=k == 1, color=HEAD if k == 1 else INK, fill=band)
        if not rows:
            _put(s, 3, 1, "尚未撰写（在 analysis.json 填写 findings，或对创新点给出 verdict/summary）。", color=GREY_TXT)
            s.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(heads))
        _widths(s, [26, 56, 48, 34])
        s.freeze_panes = "B3"
        _autoheight(s, 3, max(3, 2 + len(rows)))
        _finish(s)

    # ---------- 维度对比（多产品，产品为列）----------
    def sh_dimension():
        s = wb.create_sheet(sheet_name["维度对比"])
        _title(s, "维度对比（产品为列）", "按“基本信息 → 投保与核保准入 → 续保机制 → 既往症与赔付规则 → 核心额度 → 特殊责任 → 保费 → 定位”排列；浅青底色表示取值不同。",
               len(prods) + 1)
        _header(s, 4, ["对比维度"] + names)
        r = 5
        for sec in sections:
            _put(s, r, 1, sec, bold=True, fill=SECTION, color=NAVY)
            for j in range(len(prods)):
                _put(s, r, j + 2, "", fill=SECTION)
            s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(prods) + 1)
            r += 1
            for f in [x for x in schema_fields() if x["section"] == sec]:
                vals = [rec["fields"][f["key"]] for rec in recs]
                filled = {re.sub(r"\s+", "", v.get("value", "")) for v in vals if v.get("value")}
                _put(s, r, 1, f["label"], bold=True, fill=KEY if len(filled) > 1 else LABEL)
                for k in range(len(prods)):
                    fcell(s, r, k + 2, k, f["key"])
                r += 1
        _widths(s, [22] + [42] * len(prods))
        s.freeze_panes = "B5"
        _autoheight(s, 5, r - 1)
        _finish(s)

    # ---------- 保障责任明细 ----------
    def sh_benefits():
        s = wb.create_sheet(sheet_name["保障责任明细"])
        _title(s, "保障责任明细（按计划展开）", "每项责任拆为 限额 / 免赔额 / 赔付比例 / 范围与说明；“—”表示该计划不含此责任或未取得。", 8)
        cols = []  # (pid_idx, plan)
        for k, rec in enumerate(recs):
            plans = []
            for row in rec.get("benefits") or []:
                pl = row.get("plan") or "全部计划"
                if pl not in plans:
                    plans.append(pl)
            for pl in plans or ["—"]:
                cols.append((k, pl))
        _put(s, 4, 1, "保险责任", bold=True, fill=NAVY, color="FFFFFF", center=True)
        _put(s, 4, 2, "项目", bold=True, fill=NAVY, color="FFFFFF", center=True)
        _put(s, 5, 1, "", fill=NAVY)
        _put(s, 5, 2, "", fill=NAVY)
        s.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
        s.merge_cells(start_row=4, start_column=2, end_row=5, end_column=2)
        j0 = 3
        for k in range(len(recs)):
            n = sum(1 for kk, _ in cols if kk == k)
            for jj in range(n):
                _put(s, 4, j0 + jj, names[k] if jj == 0 else "", bold=True, fill=TEAL, color="FFFFFF", center=True)
            if n > 1:
                s.merge_cells(start_row=4, start_column=j0, end_row=4, end_column=j0 + n - 1)
            j0 += n
        for j, (k, pl) in enumerate(cols, 3):
            _put(s, 5, j, pl, bold=True, fill=SECTION, color=NAVY, center=True)
        resp_names = []
        for rec in recs:
            for row in rec.get("benefits") or []:
                if row["name"] not in resp_names:
                    resp_names.append(row["name"])
        r = 6
        for rn in resp_names:
            start = r
            for item, key in (("限额", "limit"), ("免赔额", "deductible"), ("赔付比例", "ratio"), ("范围 / 说明", "scope")):
                _put(s, r, 2, item, bold=True, fill=LABEL, center=True)
                for j, (k, pl) in enumerate(cols, 3):
                    hit = next((x for x in recs[k].get("benefits") or [] if x["name"] == rn and (x.get("plan") or "全部计划") == pl), None)
                    if hit is None:
                        hit = next((x for x in recs[k].get("benefits") or [] if x["name"] == rn and not x.get("plan")), None)
                    v = (hit.get(key) or "—") if hit else "—"
                    if hit and key == "scope":
                        ri = (recs[k].get("benefits") or []).index(hit)
                        ids, i = [], 0
                        while (recs[k]["pid"], "benefit", ri, i) in ev_idx:
                            ids.append(ev_idx[(recs[k]["pid"], "benefit", ri, i)])
                            i += 1
                        v = f"{v} {fmt_refs(ids)}".strip()
                    c = _put(s, r, j, v, center=key != "scope")
                    if hit and hit.get("status") == "待复核":
                        _style_status(c, "待复核")
                r += 1
            _put(s, start, 1, rn, bold=True, fill=LABEL, center=True)
            s.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        if not resp_names:
            _put(s, 6, 1, "尚未录入保障责任明细（record apply 写入 benefits）", color=GREY_TXT)
        _widths(s, [22, 12] + [30] * len(cols))
        s.freeze_panes = "C6"
        _autoheight(s, 6, max(6, r - 1))
        _finish(s)

    # ---------- 保费对比 ----------
    def sh_premium():
        s = wb.create_sheet(sheet_name["保费对比"])
        _title(s, "保费对比（单位：元）", "上表为 50 岁保费口径汇总；下方按费率口径分块列示，各列逐行引自费率表原件。", 8)
        _header(s, 4, ["产品", "50岁男性保费（元/年）", "保费口径与版本", "费率因子与折扣"])
        for i, (p, rec) in enumerate(zip(prods, recs), 5):
            _put(s, i, 1, names[i - 5].replace("\n", " · "), bold=True)
            for j, key in enumerate(("premium_50m", "premium_basis", "rate_factors"), 2):
                fcell(s, i, j, i - 5, key)
        r = 6 + len(prods)
        bases: dict[str, list] = {}
        for k, rec in enumerate(recs):
            for t_i, t in enumerate(rec.get("rate_tables") or []):
                bases.setdefault(re.sub(r"\s+", "", t.get("basis", "")), []).append((k, t_i, t))
        for b_i, (basis, tables) in enumerate(bases.items()):
            _put(s, r, 1, f"{chr(65 + b_i)}. {tables[0][2].get('basis')}", bold=True, fill=SECTION, color=NAVY)
            r += 1
            headers = ["投保年龄"]
            for k, t_i, t in tables:
                eid = ev_idx.get((recs[k]["pid"], "rate", t_i), "")
                for c in t["columns"]:
                    headers.append(f"{prods[k]['product_name']}\n{c}" + (f" [{eid}]" if eid else ""))
            _header(s, r, headers, fill=TEAL)
            r += 1
            ages = []
            for _, _, t in tables:
                for row in t["rows"]:
                    if row["age"] not in ages:
                        ages.append(row["age"])

            def age_key(a):
                m = re.search(r"\d+", a)
                return (0 if re.search(r"天|日", a) else 1, int(m.group()) if m else 999)
            ages.sort(key=age_key)
            for a in ages:
                hl = bool(re.search(r"46\D+50", a))
                _put(s, r, 1, a, bold=True, fill=KEY if hl else LABEL, center=True)
                col = 2
                for _, _, t in tables:
                    row = next((x for x in t["rows"] if x["age"] == a), None)
                    for ci in range(len(t["columns"])):
                        _put(s, r, col, (row["values"][ci] if row and ci < len(row["values"]) else "—"), center=True, fill=KEY if hl else None)
                        col += 1
                r += 1
            r += 1
        _widths(s, [26, 30, 40, 36] + [16] * 30)
        _autoheight(s, 5, 4 + len(prods))
        _finish(s)

    # ---------- 健告与既往症 ----------
    def sh_hd():
        s = wb.create_sheet(sheet_name["健告与既往症"])
        _title(s, "健康告知与既往症处理", "投保端（能否投保）与理赔端（是否除外、首年比例）分开列示；健告尽量逐字。", len(prods) + 1)
        _header(s, 4, ["项目"] + names)
        groups = [("投保端：健康告知与准入", ["body_type", "access_channel", "hd_count", "hd_text", "hd_disease_scope", "hd_history_items",
                                       "hd_exemptions", "hd_disclosure_effect", "nodules_three_highs", "underwriting_tools", "other_access"]),
                  ("理赔端：既往症与等待期", ["preexisting_definition", "preexisting_general", "preexisting_serious", "preexisting_special",
                                       "first_year_rules", "waiting_period"])]
        r = 5
        for g, keys in groups:
            _put(s, r, 1, g, bold=True, fill=SECTION, color=NAVY)
            for j in range(len(prods)):
                _put(s, r, j + 2, "", fill=SECTION)
            s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(prods) + 1)
            r += 1
            for key in keys:
                _put(s, r, 1, fm[key]["label"], bold=True, fill=LABEL, center=True)
                for k in range(len(prods)):
                    fcell(s, r, k + 2, k, key)
                r += 1
        _widths(s, [22] + [56] * len(prods))
        s.freeze_panes = "B5"
        _autoheight(s, 5, r - 1)
        _finish(s)

    # ---------- 版本对比 ----------
    def sh_versions():
        s = wb.create_sheet(sheet_name["版本对比"])
        _title(s, "新旧版本对比", "方向为候选判断（数值化字段），须结合条款复核；“研究者解读”来自 analysis.json。", 9)
        r = 4
        notes = analysis.get("version_notes") or {}
        for v in cmp.get("versions", []):
            o, n = v["old"], v["new"]
            _put(s, r, 1, f"{o['product_name']}（{o.get('launch') or '上市时间待核'}）  →  {n['product_name']}（{n.get('launch') or '上市时间待核'}）",
                 bold=True, fill=SECTION, color=NAVY)
            s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
            r += 1
            note = notes.get(f"{o['pid']}->{n['pid']}") or notes.get(f"{o['product_name']}->{n['product_name']}")
            if note:
                _put(s, r, 1, "研究者解读", bold=True, fill=LABEL)
                _put(s, r, 2, text(note if isinstance(note, str) else note.get("summary", "")))
                s.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
                r += 1
            _header(s, r, ["维度", "旧版", "新版", "变化", "方向（候选）", "新增要点", "删除要点", "旧版依据", "新版依据"], fill=TEAL)
            r += 1
            for ch in v["changes"]:
                orefs = fmt_refs(refs_for(ev_idx, o["pid"], ch["field"]))
                nrefs = fmt_refs(refs_for(ev_idx, n["pid"], ch["field"]))
                vals = [ch["label"], f"{ch['old']} {orefs}".strip(), f"{ch['new']} {nrefs}".strip(), ch["change"], ch["direction"],
                        "\n".join(ch["added"]), "\n".join(ch["removed"]),
                        "\n".join(f"{e.get('doc')} 第{e.get('page')}页" for e in ch["old_evidence"]),
                        "\n".join(f"{e.get('doc')} 第{e.get('page')}页" for e in ch["new_evidence"])]
                for j, x in enumerate(vals, 1):
                    c = _put(s, r, j, x, bold=j == 1, fill=LABEL if j == 1 else None)
                    if j == 5 and x:
                        c.font = _font(True, "C00000")
                r += 1
            r += 1
        if not cmp.get("versions"):
            _put(s, 4, 1, "未发现同系列多版本（或尚未运行 compare）。可用 compare --pair 旧:新 指定。", color=GREY_TXT)
        _widths(s, [20, 40, 40, 10, 24, 30, 30, 22, 22])
        s.freeze_panes = "B4"
        _autoheight(s, 4, max(4, r))
        _finish(s)

    # ---------- 创新点识别 ----------
    def sh_innov():
        s = wb.create_sheet(sheet_name["创新点识别"])
        base = inn.get("baseline", {})
        _title(s, "创新点识别", f"市场基线：样本摘要目录 {base.get('builtin_samples', 0)} 款 + 本次研究 {base.get('workspace_products', 0)} 款。"
               "判定为候选，“首见”仅指样本内；研究者结论来自 analysis.json。", 11)
        _header(s, 4, ["产品", "特征", "类别", "候选判定", "研究者结论", "本系列新增", "样本普及度", "同业示例（最早在前）", "产品原文依据", "来源", "结论边界/提示"])
        r = 5
        judg = analysis.get("innovations") or {}
        for pid, info in (inn.get("products") or {}).items():
            jl = {x.get("feature"): x for x in judg.get(pid, [])} if isinstance(judg.get(pid), list) else {}
            for c in info["candidates"]:
                j = jl.get(c["feature"], {})
                ev = c.get("evidence") or []
                vals = [f"{info['company']}·{info['product']}", c["label"], c["category"], c["verdict"],
                        (f"{j.get('verdict', '')}：{text(j.get('summary', ''))}".strip("：") if j else ""),
                        {True: "是", False: "否", None: "—"}[c.get("new_in_series")],
                        f"{c['peers_with']}/{c['peers_total']}",
                        "\n".join(f"{x['name']}（{x['launch'] or '时间待核'}）" for x in c["peer_examples"][:3]) or "无",
                        "\n".join(f"「{truncate(x.get('quote') or '', 80)}」" for x in ev[:2]) or c.get("product_snippet", ""),
                        "\n".join(f"{x.get('doc')} 第{x.get('page')}页 证据{x.get('evidence_level')}" for x in ev[:2]),
                        "；".join(c.get("flags", []) + ([j["boundary"]] if j.get("boundary") else []))]
                for k, x in enumerate(vals, 1):
                    cell = _put(s, r, k, x)
                    if k == 4 and c["verdict"] != "行业常见/跟进":
                        cell.font = _font(True, "C00000")
                r += 1
        if r == 5:
            _put(s, 5, 1, "尚未运行 innovate", color=GREY_TXT)
        _widths(s, [26, 22, 10, 12, 40, 8, 10, 34, 46, 26, 34])
        s.freeze_panes = "C5"
        _autoheight(s, 5, max(5, r - 1))
        _finish(s)

    # ---------- 证据索引 ----------
    def sh_evidence():
        s = wb.create_sheet(sheet_name["证据索引"])
        _title(s, "证据索引", "每条引文均已在归档原件的抽取文本中核验（匹配=exact 逐字 / fuzzy 近似）；文件列可直接打开归档原件。", 12)
        _header(s, 4, ["编号", "产品", "字段", "取值", "状态", "原文引文", "文件", "页码", "匹配", "证据等级", "来源链接"])
        for i, e in enumerate(ev_items, 5):
            file_link = Path(ws_.root / e["archive_dir"] / (e.get("doc") or "")).as_uri() if e.get("doc") else None
            vals = [e["id"], f"{e['company']}·{e['product']}", e["label"], e["value"], e["status"], e.get("quote"), e.get("doc"),
                    e.get("page"), f"{e.get('match')}{'(' + str(e.get('score')) + ')' if e.get('match') == 'fuzzy' else ''}",
                    e.get("evidence_level"), e.get("url")]
            for k, x in enumerate(vals, 1):
                _put(s, i, k, x, center=k in (1, 5, 8, 9, 10), link=(file_link if k == 7 else (x if k == 11 and x else None)))
        _widths(s, [7, 24, 18, 34, 10, 60, 30, 6, 10, 8, 40])
        s.freeze_panes = "C5"
        s.auto_filter.ref = f"A4:K{max(5, 4 + len(ev_items))}"
        _autoheight(s, 5, 4 + len(ev_items))
        _finish(s)

    # ---------- 资料清单 ----------
    def sh_docs():
        s = wb.create_sheet(sheet_name["资料清单"])
        _title(s, "资料清单与完整度", "SHA256 可用于复核文件未被改动；版本核对=正文是否出现已登记备案条款名。", 11)
        _header(s, 4, ["产品", "类型", "文件", "页数", "证据等级", "版本核对", "SHA256", "来源链接", "获取时间", "提示"])
        r = 5
        for p in prods:
            m = load_manifest(ws_, p["pid"])
            for d in m["documents"]:
                src = (d.get("sources") or [{}])[0]
                vals = [f"{p['company']}·{p['product_name']}", f"{d['doc_type']} {d['doc_type_name']}", d["file"], d.get("pages"),
                        d.get("evidence_level"), d.get("version_check", {}).get("status"), d.get("sha256"), src.get("url") or src.get("local_path", ""),
                        src.get("retrieved_at", ""), d.get("warning", "") or d.get("version_check", {}).get("year_note", "")]
                link = Path(ws_.root / p["archive_dir"] / d["file"]).as_uri()
                for k, x in enumerate(vals, 1):
                    _put(s, r, k, x, center=k in (4, 5, 6), link=(link if k == 3 else (x if k == 8 and str(x).startswith("http") else None)))
                r += 1
            for g in m["gaps"]:
                vals = [f"{p['company']}·{p['product_name']}", f"{g['doc_type']} {g['doc_type_name']}", g.get("file", ""), "", "", g["status"], "", "",
                        g.get("marked_at", ""), g.get("note", "")]
                for k, x in enumerate(vals, 1):
                    _put(s, r, k, x, color=GREY_TXT)
                r += 1
        _widths(s, [26, 18, 44, 6, 8, 10, 20, 44, 20, 36])
        s.freeze_panes = "B5"
        _autoheight(s, 5, max(5, r - 1))
        _finish(s)

    # ---------- 校验结果 ----------
    def sh_validation():
        s = wb.create_sheet(sheet_name["校验结果"])
        _title(s, "校验结果", f"生成于 {val.get('generated_at', '未运行 validate')}；error 必须清零后才能把相关内容写入结论。", 3)
        _header(s, 4, ["产品", "级别", "问题"])
        r = 5
        for pr in val.get("products", []):
            for lvl, lst in (("error", pr["errors"]), ("warning", pr["warnings"])):
                for msg in lst:
                    _put(s, r, 1, f"{pr['company']}·{pr['product']}")
                    _put(s, r, 2, lvl, center=True, color="C00000" if lvl == "error" else "7F6000", bold=lvl == "error")
                    _put(s, r, 3, msg)
                    r += 1
        if r == 5:
            _put(s, 5, 1, "无问题或尚未运行 validate", color=GREY_TXT)
        _widths(s, [30, 10, 120])
        _autoheight(s, 5, max(5, r - 1))
        _finish(s)

    builders = {"说明": sh_readme, "系列总览": sh_series_overview, "产品总览": sh_market_overview, "严重既往症对比": sh_preexisting,
                "产品演进与趋势": sh_evolution, "核心发现": sh_findings, "维度对比": sh_dimension, "保障责任明细": sh_benefits,
                "保费对比": sh_premium, "健告与既往症": sh_hd, "版本对比": sh_versions, "创新点识别": sh_innov, "证据索引": sh_evidence,
                "资料清单": sh_docs, "校验结果": sh_validation}
    for n in order:
        builders[n]()
    wb.active = 1

    out = out or (ws_.outputs_dir / f"{title}_对比表_{today().replace('-', '')}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
    except PermissionError:
        alt = out.with_name(out.stem + "_new" + out.suffix)
        wb.save(alt)
        out = alt
    return out


SHEET_DESC = {
    "系列总览": "维度 × 版本的全字段对比，含责任覆盖速查与保费速览；浅青底色为版本间差异",
    "产品总览": "每行一款产品的核心形态（基本信息 / 投保与健康告知 / 既往症与理赔 / 核心额度与特殊责任 / 保费），可筛选",
    "严重既往症对比": "既往症定义、一般/严重既往症处理与首年规则，附严重既往症病种速查",
    "产品演进与趋势": "同系列新旧版本在健告、既往症与赔付、保障责任、续保与保费上的变化及研究者解读",
    "核心发现": "研究者归纳的核心发现、风险控制逻辑与结论边界",
    "维度对比": "维度 × 产品的全字段对比",
    "保障责任明细": "各责任按计划展开的限额 / 免赔额 / 赔付比例 / 范围",
    "保费对比": "50 岁保费口径汇总，以及逐行引自费率表的分年龄费率（46–50 岁档高亮）",
    "健告与既往症": "投保端健康告知与准入、理赔端既往症与等待期的逐字对比",
    "版本对比": "新旧版本逐字段变化、方向判断（候选）与依据页码",
    "创新点识别": "特征候选判定、样本普及度、同业示例、原文依据与研究者结论",
    "证据索引": "全部 [E#] 证据：原文引文、文件（可点击打开）、页码、匹配方式、证据等级、来源链接",
    "资料清单": "归档文件、证据等级、版本核对、SHA256、来源与资料缺口",
    "校验结果": "validate 的 error / warning 明细",
}
