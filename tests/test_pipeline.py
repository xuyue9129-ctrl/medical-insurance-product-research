"""集成测试：登记→导入→抽取→写入（引文核验）→费率→缺口→校验→对比→创新点→Excel→报告。"""
import json
import unittest
from pathlib import Path

from _util import FIXTURES, HAS_PDF, TempWorkspaceCase, has_module


@unittest.skipUnless(HAS_PDF, "需要任一 PDF 解析库")
class PipelineTest(TempWorkspaceCase):
    def setUp(self):
        super().setUp()
        from medins_lib.fetch import archive_file
        self.a = self.ws.register({"company": "示例健康", "product_name": "示例·免健告百万医疗（2025）", "series": "示例·百万医疗",
                                   "insurer_full": "示例健康保险股份有限公司", "clause_names": ["示例互联网个人住院A款医疗保险"]})
        self.b = self.ws.register({"company": "示例健康", "product_name": "示例·免健告百万医疗（2026）", "series": "示例·百万医疗",
                                   "insurer_full": "示例健康保险股份有限公司", "clause_names": ["示例互联网个人住院B款医疗保险"]})
        self.docA = archive_file(self.ws, self.a["pid"], FIXTURES / "示例A款条款.pdf",
                                 source={"url": "https://example.com/a.pdf"}, title="示例互联网个人住院A款医疗保险条款")
        self.docB = archive_file(self.ws, self.b["pid"], FIXTURES / "示例B款条款.pdf", title="示例互联网个人住院B款医疗保险条款")
        self.rateB = archive_file(self.ws, self.b["pid"], FIXTURES / "示例B款费率表.pdf", title="示例互联网个人住院B款医疗保险费率表")

    def test_archive_metadata(self):
        self.assertEqual(self.docA["doc_type"], "01")
        self.assertEqual(self.rateB["doc_type"], "02")
        self.assertEqual(self.docA["version_check"]["status"], "matched")
        self.assertEqual(len(self.docA["sha256"]), 64)
        self.assertEqual(self.docA["detected_clause_names"], ["示例互联网个人住院A款医疗保险"])
        self.assertIn("示例健康险[2025]医疗保险001号", self.docA["detected_clause_codes"])
        # 同一文件再次归档 → 识别为重复，只追加来源
        from medins_lib.fetch import archive_file, load_manifest
        dup = archive_file(self.ws, self.a["pid"], FIXTURES / "示例A款条款.pdf", source={"url": "https://mirror.example.com/a.pdf"})
        self.assertTrue(dup.get("duplicate"))
        m = load_manifest(self.ws, self.a["pid"])
        self.assertEqual(len(m["documents"]), 1)
        self.assertEqual(len(m["documents"][0]["sources"]), 2)
        inv = (self.ws.product_dir(self.a["pid"]) / "00_资料清单.txt").read_text(encoding="utf-8")
        self.assertIn("核心资料完整度", inv)

    def test_extract_safe_fields(self):
        from medins_lib.extract import run_extract
        from medins_lib.records import load_record
        run_extract(self.ws, self.b["pid"], apply_safe=True)
        F = load_record(self.ws, self.b["pid"])["fields"]
        self.assertEqual(F["entry_age"]["value"], "出生满30天至80周岁")
        self.assertEqual(F["term"]["value"], "1年")
        self.assertEqual(F["guaranteed_renewal"]["value"], "不保证续保")
        self.assertEqual(F["waiting_period"]["value"], "疾病等待期30日")
        self.assertEqual(F["insurer_full"]["value"], "示例健康保险股份有限公司")
        self.assertEqual(F["clause_codes"]["value"], "示例健康险[2026]医疗保险009号")
        for k in ("entry_age", "term", "waiting_period"):
            self.assertEqual(F[k]["status"], "规则抽取")
            self.assertEqual(F[k]["evidence"][0]["match"], "exact")

    def test_record_apply_verifies_quotes(self):
        from medins_lib.records import apply_patch, load_record
        res = apply_patch(self.ws, self.b["pid"], [
            {"field": "deductible", "value": "一般医疗年免赔额1.5万元",
             "evidence": [{"doc": "01_", "quote": "本合同一般医疗保险金的年免赔额为1.5万元。"}]},
            {"field": "coinsurance", "value": "一般医疗100%",
             "evidence": [{"doc": "条款", "quote": "本合同保证续保20年，赔付比例100%"}]},          # 编造的引文
            {"field": "免赔额", "value": "一般医疗年免赔额2万元",
             "evidence": [{"doc": "01_", "quote": "本合同一般医疗保险金的年免赔额为1.5万元。"}]},  # 数字对不上
            {"field": "hd_text", "status": "待核"},                                                 # 缺原因
            {"field": "positioning", "status": "推定", "value": "门槛低的一年期免健告百万医疗", "note": "据条款归纳"},
        ])
        self.assertTrue(res[0]["applied"])
        self.assertFalse(res[1]["applied"])
        self.assertIn("引文未在原文中找到", res[1]["problems"][0])
        self.assertTrue(res[2]["applied"])
        self.assertTrue(any("数字 2" in w for w in res[2]["warnings"]))
        self.assertFalse(res[3]["applied"])
        self.assertTrue(res[4]["applied"])
        rec = load_record(self.ws, self.b["pid"])
        ev = rec["fields"]["deductible"]["evidence"][0]
        self.assertEqual((ev["match"], ev["page"]), ("exact", 1))
        self.assertEqual(ev["doc"], self.docB["file"])

    def test_rates_import_and_full_build(self):
        from medins_lib.compare import run_compare
        from medins_lib.excel import build_workbook
        from medins_lib.extract import run_extract
        from medins_lib.fetch import mark_gap
        from medins_lib.innovation import run_innovation
        from medins_lib.rates import block_to_table, parse_rate_file
        from medins_lib.records import apply_patch
        from medins_lib.report import build_report
        from medins_lib.validate import validate_all

        for p, doc in ((self.a, self.docA), (self.b, self.docB)):
            run_extract(self.ws, p["pid"], apply_safe=True)
        blocks = parse_rate_file(self.ws.root / self.rateB["text_file"])
        self.assertEqual(len(blocks[0]["rows"]), 5)  # 0-5、6-10、41-45 + 跨页 46-50、51-55
        table = block_to_table(blocks[0], ["计划一", "计划二"], [0, 1], "有社保·年交·含等待期", self.rateB["file"])
        r = apply_patch(self.ws, self.b["pid"], [{"field": "rate_table", "table": table}])
        self.assertTrue(r[0]["applied"], r)

        apply_patch(self.ws, self.a["pid"], [
            {"field": "deductible", "value": "一般医疗年免赔额1万元", "evidence": [{"doc": "01_", "quote": "本合同一般医疗保险金的年免赔额为1万元。"}]},
            {"field": "coinsurance", "value": "一般医疗赔付比例80%", "evidence": [{"doc": "01_", "quote": "一般医疗保险金的赔付比例A为80%"}]},
            {"field": "advanced_therapy", "value": "质子重离子100万元，0免赔", "evidence": [{"doc": "01_", "quote": "恶性肿瘤质子重离子医疗保险金：限额100万元，免赔额为0。"}]},
            {"field": "launch_date", "value": "2025-04", "status": "推定", "note": "测试数据"},
        ])
        apply_patch(self.ws, self.b["pid"], [
            {"field": "deductible", "value": "一般医疗年免赔额1.5万元", "evidence": [{"doc": "01_", "quote": "本合同一般医疗保险金的年免赔额为1.5万元。"}]},
            {"field": "coinsurance", "value": "一般医疗赔付比例100%", "evidence": [{"doc": "01_", "quote": "一般医疗保险金的赔付比例A为100%"}]},
            {"field": "rehab", "value": "康复住院100万元、0免赔、80%", "evidence": [{"doc": "01_", "quote": "康复住院医疗保险金：限额100万元，免赔额为0，赔付比例A为80%。"}]},
            {"field": "preexisting_general", "value": "一般既往症首年：计划一按约定比例×50%，计划二不减赔",
             "evidence": [{"doc": "01_", "quote": "计划一，我们将赔付比例A在上述约定的基础上乘以50%进行赔付；计划二不减赔。"}]},
            {"field": "premium_50m", "value": "计划一542；计划二886", "evidence": [{"doc": "费率表", "quote": "46-50 542 886 48.78 79.74"}]},
            {"field": "premium_basis", "value": "46–50岁档、有社保、年交、含等待期", "evidence": [{"doc": "费率表", "quote": "有基本医疗保险或公费医疗"}]},
            {"field": "launch_date", "value": "2026-06", "status": "推定", "note": "测试数据"},
            {"field": "benefits", "rows": [{"name": "一般医疗保险金", "plan": "计划一", "deductible": "1.5万元", "ratio": "100%",
                                            "evidence": [{"doc": "01_", "quote": "本合同一般医疗保险金的年免赔额为1.5万元。"}]}]},
        ])
        mark_gap(self.ws, self.b["pid"], "03", "无健康告知（非缺文件）", note="虚构产品：免健告")

        rep = validate_all(self.ws)
        self.assertGreater(rep["totals"]["errors"], 0)  # 必填字段未全部处理，应报 error
        msgs = " ".join(rep["products"][1]["errors"])
        self.assertNotIn("03 健康告知 缺失", msgs)

        cmp = run_compare(self.ws)
        self.assertEqual(len(cmp["versions"]), 1)
        ch = {c["field"]: c for c in cmp["versions"][0]["changes"]}
        self.assertIn("免赔额提高", ch["deductible"]["direction"])
        self.assertIn("赔付比例提高", ch["coinsurance"]["direction"])
        self.assertIn("出生满30天至80周岁", ch["entry_age"]["new"])

        # 本机样本摘要目录（虚构）：作为创新点识别的市场基线
        import os
        from medins_lib.common import write_json
        write_json(Path(os.environ["MEDINS_HOME"]) / "catalog.json", [
            {"company": "虚构财险", "product_name": "虚构百万医疗2025", "launch": "2025-03", "overview_as_of": "2026-01-01",
             "overview": {"rehab": "康复医疗10万，0免赔", "access_channel": "免健告", "advanced_therapy": "质子重离子"}},
            {"company": "虚构健康", "product_name": "虚构中端医疗2026", "launch": "2026-02", "overview_as_of": "2026-01-01",
             "overview": {"rehab": "—", "access_channel": "简易健告", "preexisting_special": "高危结节专项限额1万"}},
        ])
        inn = run_innovation(self.ws)
        feats = {c["feature"]: c for c in inn["products"][self.b["pid"]]["candidates"]}
        self.assertTrue(feats["rehab"]["new_in_series"])
        self.assertIn("pre_first_year_option", feats)
        self.assertEqual(inn["baseline"]["builtin_samples"], 2)
        self.assertEqual(feats["rehab"]["peers_with"], 1)  # 虚构百万医疗2025 有康复；“—”不算

        (self.ws.analysis_path).write_text(json.dumps({
            "version_notes": {f"{self.a['pid']}->{self.b['pid']}": f"免赔额由1万元提高到1.5万元[[{self.b['pid']}:deductible]]。"},
            "innovations": {self.b["pid"]: [{"feature": "pre_first_year_option", "label": "一般既往症首年比例可选", "verdict": "差异化",
                                             "summary": f"计划一首年×50%[[{self.b['pid']}:preexisting_general]]", "boundary": "样本内"}]},
            "overall": {"summary": "免赔额上调到1.5万元。"}}, ensure_ascii=False), encoding="utf-8")

        if has_module("openpyxl"):
            import openpyxl
            # 同系列两版本 → 系列版式
            wb = openpyxl.load_workbook(build_workbook(self.ws))
            self.assertEqual(wb.sheetnames[:3], ["说明", "① 系列总览", "② 保障责任明细"])
            ev_sheet = wb["⑦ 证据索引"]
            self.assertEqual(ev_sheet.cell(5, 1).value, "E1")
            self.assertGreater(ev_sheet.max_row, 10)
            hd_labels = [wb["④ 健告与既往症"].cell(r, 1).value for r in range(5, 30)]
            for lab in ("适用人群定位（次标体/健康体）", "健告疾病问询范围", "住院/手术/检查异常问项（含时限）",
                        "可正常投保的情形", "告知效力与问询外疾病处理"):
                self.assertIn(lab, hd_labels)
            prem = wb["③ 保费对比"]
            values = [c.value for row in prem.iter_rows() for c in row if c.value]
            self.assertIn("542", values)
            ov = [c.value for row in wb["① 系列总览"].iter_rows() for c in row if c.value]
            self.assertTrue(any(str(v).startswith("责任覆盖速查") for v in ov))
            self.assertIn("✔", ov)  # 康复保障已写入 → ✔
            vtxt = " ".join(str(c.value) for row in wb["⑤ 版本对比"].iter_rows() for c in row if c.value)
            self.assertNotIn("[[", vtxt)  # 研究者解读中的引用记号已换成 [E#]
            # 多产品横比版式
            wb = openpyxl.load_workbook(build_workbook(self.ws, mode="market", out=self.ws.outputs_dir / "market.xlsx"))
            self.assertEqual(wb.sheetnames[:5], ["说明", "① 产品总览", "② 严重既往症对比", "③ 产品演进与趋势", "④ 核心发现"])
            s = wb["① 产品总览"]
            self.assertEqual(s.freeze_panes, "D4")
            self.assertEqual(s.cell(3, 1).value, "序号")
            self.assertEqual(s.cell(4, 1).value, 1)
            evo = " ".join(str(c.value) for row in wb["③ 产品演进与趋势"].iter_rows() for c in row if c.value)
            self.assertIn("免赔额", evo)
            finds = " ".join(str(c.value) for row in wb["④ 核心发现"].iter_rows() for c in row if c.value)
            self.assertIn("一般既往症首年比例可选（差异化）", finds)
        res = build_report(self.ws, docx=has_module("docx"))
        md = Path(res["markdown"]).read_text(encoding="utf-8")
        self.assertIn("[E", md)
        self.assertIn("一般既往症首年比例可选", md)
        self.assertNotIn("[[", md)  # 引用记号已全部替换
        self.assertIn("## 二、承保范围", md)
        self.assertRegex(md, r"\n表1　")
        self.assertTrue(any("overall.summary" in p for p in res["problems"]))  # 含数字但无引用的句子被提示
        if has_module("docx"):
            import docx
            self.assertTrue(res["docx"].endswith(".docx"))
            d = docx.Document(res["docx"])
            styles = {p.style.name for p in d.paragraphs}
            for st in ("Report Title CN", "Heading 1", "Heading 2", "Table Caption CN", "Body Text CN", "TOC Title CN"):
                self.assertIn(st, styles)
            self.assertGreaterEqual(len(d.tables), 5)
        if has_module("pptx"):
            from pptx import Presentation
            from medins_lib.ppt import build_ppt
            r = build_ppt(self.ws, template="")
            prs = Presentation(r["pptx"])
            self.assertEqual(len(prs.slides), r["slides"])
            self.assertGreaterEqual(r["slides"], 10)
            texts = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
            self.assertIn("核保端：健康告知与核保准入", texts)
            self.assertIn("一般既往症首年比例可选", texts)
            self.assertNotIn("[[", texts)


if __name__ == "__main__":
    unittest.main()
